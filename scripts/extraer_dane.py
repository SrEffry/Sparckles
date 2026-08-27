import openpyxl, json, unicodedata, collections

wb = openpyxl.load_workbook('exo.xlsx', read_only=True, data_only=True)
sh = wb[[s for s in wb.sheetnames if 'digos' in s][0]]
rows = list(sh.iter_rows(values_only=True))[4:]   # las 4 primeras son título/encabezado

def txt(v):
    return unicodedata.normalize('NFC', str(v).strip()) if v is not None else None

deptos, municipios, paises = [], [], []
for r in rows:
    r = list(r) + [None] * (8 - len(r))
    if txt(r[0]) and txt(r[1]):
        deptos.append({"codigo": txt(r[0]).zfill(2), "nombre": txt(r[1])})
    if txt(r[3]) and txt(r[4]):
        municipios.append({"codigo": txt(r[4]).zfill(5), "nombre": txt(r[3])})
    if txt(r[6]) and txt(r[7]):
        paises.append({"codigo": txt(r[6]).zfill(3), "nombre": txt(r[7])})

# ---- Validación: sin esto no vale la pena guardarlo ----
fallas = []
cod_dep = {d["codigo"] for d in deptos}
if len(cod_dep) != len(deptos): fallas.append("departamentos duplicados")

vistos = collections.Counter(m["codigo"] for m in municipios)
dup = [c for c, n in vistos.items() if n > 1]
if dup: fallas.append(f"municipios duplicados: {dup[:5]}")

malos = [m for m in municipios if len(m["codigo"]) != 5 or not m["codigo"].isdigit()]
if malos: fallas.append(f"municipios con código no numérico de 5: {malos[:5]}")

huerfanos = [m for m in municipios if m["codigo"][:2] not in cod_dep]
if huerfanos: fallas.append(f"municipios cuyo prefijo no es un departamento conocido: {huerfanos[:5]}")

malp = [p for p in paises if len(p["codigo"]) != 3 or not p["codigo"].isdigit()]
if malp: fallas.append(f"países con código raro: {malp[:5]}")

print("departamentos:", len(deptos), "| municipios:", len(municipios), "| países:", len(paises))
print("Colombia:", [p for p in paises if 'COLOMBIA' in p['nombre'].upper()])
print("muestra municipios:", municipios[:3], "...", municipios[-2:])
por_dep = collections.Counter(m["codigo"][:2] for m in municipios)
print("municipios por departamento (top 5):", por_dep.most_common(5))
print("departamentos sin municipios:", sorted(cod_dep - set(por_dep)))
print("\nFALLAS:", fallas if fallas else "ninguna")

if not fallas:
    # Cada municipio lleva su departamento resuelto: el layout pide los dos códigos por
    # separado y derivarlo en cada consulta sería repetir la partición en veinte sitios.
    nombre_dep = {d["codigo"]: d["nombre"] for d in deptos}
    for m in municipios:
        m["departamento"] = m["codigo"][:2]
        m["nombreDepartamento"] = nombre_dep[m["codigo"][:2]]
    salida = {
        "fuente": "Hoja 'Códigos dtos, mun, paises' de 'Formato Exógena 2025 Formatos informantes basicos.xlsx' (Wiliam Dussán Salazar)",
        "departamentos": sorted(deptos, key=lambda d: d["codigo"]),
        "municipios": sorted(municipios, key=lambda m: m["codigo"]),
        "paises": sorted(paises, key=lambda p: p["codigo"]),
    }
    with open('dane.json', 'w', encoding='utf-8') as f:
        json.dump(salida, f, ensure_ascii=False, indent=1)
    print("escrito dane.json")
